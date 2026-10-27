from datetime import datetime
from types import SimpleNamespace

import numpy as np
import soundfile as sf
from docx import Document

from meeting_notes.config import Settings
from meeting_notes.exporter import MeetingInfo, render_markdown, write_docx
from meeting_notes.pipeline import process_recording
from meeting_notes.recorder import mix_tracks
from meeting_notes.summarizer import ActionItem, MeetingSummary, Summarizer
from meeting_notes.transcriber import Segment, format_timestamp

SEGMENTS = [Segment(0, 4, "Oke kita mulai meeting-nya."), Segment(65, 70, "Budi follow up ke vendor Jumat.")]
SUMMARY = MeetingSummary(
    title="Sinkronisasi Vendor",
    summary="Tim membahas progres vendor.",
    participants=["Budi", "Sari"],
    key_points=["Progres vendor terlambat"],
    decisions=["Ganti jadwal rilis"],
    action_items=[ActionItem(task="Follow up ke vendor", owner="Budi", deadline="Jumat")],
    open_questions=[],
)
INFO = MeetingInfo(started_at=datetime(2026, 10, 27, 9, 30), duration_seconds=3725)


def test_format_timestamp():
    assert format_timestamp(65) == "01:05"
    assert format_timestamp(3725) == "01:02:05"


def test_render_markdown_with_summary():
    md = render_markdown(INFO, SEGMENTS, SUMMARY)
    assert md.startswith("# Sinkronisasi Vendor")
    assert "**Durasi:** 01:02:05" in md
    assert "| Follow up ke vendor | Budi | Jumat |" in md
    assert "## Pertanyaan Terbuka\n\n- -" in md
    assert "**[01:05]** Budi follow up ke vendor Jumat." in md


def test_render_markdown_transcript_only():
    md = render_markdown(INFO, SEGMENTS, None, note="Tanpa ringkasan")
    assert md.startswith("# Rapat 27-10-2026")
    assert "> ⚠️ Tanpa ringkasan" in md
    assert "## Ringkasan" not in md


def test_write_docx(tmp_path):
    path = write_docx(tmp_path / "notes.docx", INFO, SEGMENTS, SUMMARY)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Sinkronisasi Vendor" in text and "Transkrip Lengkap" in text
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["Follow up ke vendor", "Budi", "Jumat"]


def test_mix_tracks_handles_different_lengths(tmp_path):
    sr = 16000
    sf.write(tmp_path / "a.wav", np.full(sr, 0.25, dtype="float32"), sr, subtype="PCM_16")
    sf.write(tmp_path / "b.wav", np.full(sr // 2, 0.25, dtype="float32"), sr, subtype="PCM_16")
    out = mix_tracks([tmp_path / "a.wav", tmp_path / "b.wav"], tmp_path / "mix.wav", sr, blocksize=3000)
    data, _ = sf.read(out)
    assert len(data) == sr
    assert np.allclose(data[: sr // 2], 0.5, atol=1e-3)
    assert np.allclose(data[sr // 2 :], 0.25, atol=1e-3)


class FakeTranscriber:
    def transcribe(self, audio_path, on_progress=None):
        return SEGMENTS


def _audio(tmp_path):
    path = tmp_path / "meeting.wav"
    sf.write(path, np.zeros(16000, dtype="float32"), 16000)
    return path


def test_pipeline_without_api_key_saves_transcript(tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
    settings = Settings(output_dir=tmp_path / "out")
    result = process_recording(_audio(tmp_path), settings, INFO.started_at, lambda _: None, FakeTranscriber())
    assert result.markdown_path.name == "2026-10-27_0930_rapat.md"
    assert result.docx_path.exists()
    assert "ANTHROPIC_API_KEY" in result.note


def test_pipeline_with_summary(tmp_path, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    monkeypatch.setattr(Summarizer, "summarize", lambda self, transcript: SUMMARY)
    settings = Settings(output_dir=tmp_path / "out")
    result = process_recording(_audio(tmp_path), settings, INFO.started_at, lambda _: None, FakeTranscriber())
    assert result.markdown_path.name == "2026-10-27_0930_sinkronisasi-vendor.md"
    assert result.note is None


def test_summarizer_request(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    calls = {}

    def fake_parse(**kwargs):
        calls.update(kwargs)
        return SimpleNamespace(stop_reason="end_turn", parsed_output=SUMMARY)

    summarizer = Summarizer()
    monkeypatch.setattr(summarizer.client.beta.messages, "parse", fake_parse)
    assert summarizer.summarize("[00:00] halo") is SUMMARY
    assert calls["model"] == "claude-opus-5-5"
    assert calls["output_format"] is MeetingSummary
    assert "<transkrip>\n[00:00] halo\n</transkrip>" in calls["messages"][0]["content"]
