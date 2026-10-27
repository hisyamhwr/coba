"""Menyimpan notulen rapat sebagai Markdown (.md) dan Word (.docx)."""

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document

from .summarizer import MeetingSummary
from .transcriber import Segment, format_timestamp


@dataclass
class MeetingInfo:
    started_at: datetime
    duration_seconds: float


def _bullet_sections(summary: MeetingSummary) -> list[tuple[str, list[str]]]:
    return [
        ("Peserta", summary.participants),
        ("Poin Pembahasan", summary.key_points),
        ("Keputusan", summary.decisions),
    ]


def _title(info: MeetingInfo, summary: MeetingSummary | None) -> str:
    return summary.title if summary else f"Rapat {info.started_at:%d-%m-%Y}"


def render_markdown(
    info: MeetingInfo, segments: list[Segment], summary: MeetingSummary | None, note: str | None = None
) -> str:
    lines = [
        f"# {_title(info, summary)}",
        "",
        f"**Tanggal:** {info.started_at:%d-%m-%Y %H:%M}  ",
        f"**Durasi:** {format_timestamp(info.duration_seconds)}",
        "",
    ]
    if note:
        lines += [f"> ⚠️ {note}", ""]

    if summary:
        lines += ["## Ringkasan", "", summary.summary, ""]
        for heading, items in _bullet_sections(summary):
            lines += [f"## {heading}", ""]
            lines += [f"- {item}" for item in items] or ["- -"]
            lines.append("")
        lines += ["## Action Items", ""]
        if summary.action_items:
            lines += ["| Tugas | Penanggung Jawab | Tenggat |", "|---|---|---|"]
            lines += [
                f"| {a.task} | {a.owner} | {a.deadline} |".replace("\n", " ") for a in summary.action_items
            ]
        else:
            lines.append("- -")
        lines.append("")
        lines += ["## Pertanyaan Terbuka", ""]
        lines += [f"- {item}" for item in summary.open_questions] or ["- -"]
        lines.append("")

    lines += ["## Transkrip Lengkap", ""]
    lines += [f"**[{format_timestamp(s.start)}]** {s.text}  " for s in segments] or ["_(tidak ada suara terdeteksi)_"]
    return "\n".join(lines) + "\n"


def write_docx(
    path: Path, info: MeetingInfo, segments: list[Segment], summary: MeetingSummary | None, note: str | None = None
) -> Path:
    doc = Document()
    doc.add_heading(_title(info, summary), level=0)
    doc.add_paragraph(f"Tanggal: {info.started_at:%d-%m-%Y %H:%M}\nDurasi: {format_timestamp(info.duration_seconds)}")
    if note:
        doc.add_paragraph(f"⚠️ {note}")

    def bullets(items: list[str]) -> None:
        for item in items or ["-"]:
            doc.add_paragraph(item, style="List Bullet")

    if summary:
        doc.add_heading("Ringkasan", level=1)
        doc.add_paragraph(summary.summary)
        for heading, items in _bullet_sections(summary):
            doc.add_heading(heading, level=1)
            bullets(items)

        doc.add_heading("Action Items", level=1)
        if summary.action_items:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, ["Tugas", "Penanggung Jawab", "Tenggat"], strict=True):
                cell.text = text
                cell.paragraphs[0].runs[0].bold = True
            for a in summary.action_items:
                for cell, text in zip(table.add_row().cells, [a.task, a.owner, a.deadline], strict=True):
                    cell.text = text
        else:
            bullets([])

        doc.add_heading("Pertanyaan Terbuka", level=1)
        bullets(summary.open_questions)

    doc.add_heading("Transkrip Lengkap", level=1)
    for s in segments:
        p = doc.add_paragraph()
        p.add_run(f"[{format_timestamp(s.start)}] ").bold = True
        p.add_run(s.text)
    if not segments:
        doc.add_paragraph("(tidak ada suara terdeteksi)")

    doc.save(path)
    return path
